#!/usr/bin/env python3
"""Build the email-format Excel (Email.xlsx) from data.json + reference/keys_golive.json.
Usage: python build_xlsx.py <data.json> <out.xlsx>
Formatting replicates the user's master "Rest Activity Tracking.xlsx" Summary sheet:
Aptos Narrow 12, italic hotel/group names, E1E1DE banding, accounting number formats,
boxed main + management-rollup tables, size-16 filled Total row.
"""
import sys, json, os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
data=json.load(open(sys.argv[1])); out=sys.argv[2]
ref=json.load(open(os.path.join(os.path.dirname(__file__),"..","reference","keys_golive.json")))
asof=date.fromisoformat(data["asof"])
rorder=["CARD_DECLINED","REPEAT_OFFENSE","GUEST_COMPLAINT","OTHER","VIP","AIRLINE_GROUP","GUEST_WARNED","NO_CARD","VACANT","CHARGEBACK","CHARGED_LATER","UNBOOKED"]
rlbl={"CARD_DECLINED":"Card Declined / Insufficient Funds","REPEAT_OFFENSE":"Repeat Offense","GUEST_COMPLAINT":"Guest Complaint","OTHER":"Other","VIP":"VIP / Loyalty","AIRLINE_GROUP":"Airlines / Groups","GUEST_WARNED":"Guest Warned","NO_CARD":"No Card on File","VACANT":"Vacant / Checked Out","CHARGEBACK":"Chargeback","CHARGED_LATER":"Charged Later","UNBOOKED":"Unbooked"}
rows=data["sgl"]
GREY="E1E1DE"
FN="Aptos Narrow"
def F(sz=12,b=False,i=False,color="000000"):return Font(name=FN,size=sz,bold=b,italic=i,color=color)
RED=lambda sz=12:Font(name=FN,size=sz,bold=True,color="C00000")
GRN=lambda sz=12:Font(name=FN,size=sz,bold=True,color="00B050")
med=Side(style="medium",color="000000");thin=Side(style="thin",color="000000")
# number formats (match master)
DOL0='"$"#,##0'; PCT2='0.00%'; PCTd='0.0%'; PCT0='0%'; INT0='0'; DATE='M/D/YYYY'
ACC0='_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
ACC2='_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
ACCD='_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
HEADERS=["Hotel Name","Go Live Date","Incident / Av Room","Charge Rate","Revenue","Expense","Profit","Total Incidents","Charged Incidents","RevPAR","Capex Cost","Proj. Annual Incidents","Incremental Benefit Payback Period","Management Company"]
NC=14                                        # last visible column (Management Company)
RBASE=15; KEYC=RBASE+len(rorder); DAYC=RBASE+1+len(rorder)
KL=get_column_letter(KEYC); DL=get_column_letter(DAYC)
wb=Workbook();ws=wb.active;ws.title="Email Format";ws.sheet_view.showGridLines=False

def sides(j,top=None,bottom=None,capex_left=True):
    left=med if j==1 else (thin if (j==11 and capex_left) else None)
    right=med if j==NC else None
    return Border(top=top,bottom=bottom,left=left,right=right)

# ---- header row 1 ----
for j,h in enumerate(HEADERS,1):
    c=ws.cell(1,j,h);c.font=F(sz=12,b=True)
    c.alignment=Alignment(horizontal=("left" if j==1 else "center"),vertical="center",wrap_text=(j!=1))
    c.border=sides(j,top=med,bottom=thin,capex_left=False)
# hidden reason + keys/days headers
for k,rc in enumerate(rorder):
    c=ws.cell(1,RBASE+k,rlbl[rc]);c.font=F(sz=12,b=True);c.alignment=Alignment(horizontal="center",wrap_text=True)
ws.cell(1,KEYC,"Keys").font=F(sz=12,b=True);ws.cell(1,DAYC,"Days Since Go-Live").font=F(sz=12,b=True)
ws.row_dimensions[1].height=47.25
r0=2

def put(r,col,val,fmt=None,al="center",fill=None,font=None,border=None,italic=False):
    c=ws.cell(r,col,val)
    c.font=font or F(sz=12,i=italic)
    c.alignment=Alignment(horizontal=al,vertical="center")
    if fmt:c.number_format=fmt
    if fill:c.fill=PatternFill("solid",fgColor=fill)
    if border is not None:c.border=border
    return c

grp={}
for idx,row in enumerate(rows):
    nm=row["name"];bil=row["billable"];chg=row["charged"];net=row["net"];rc=row["rest"];R=row.get("reasons",{})
    r=r0+idx;gl=date.fromisoformat(ref[nm]["golive"]);days=(asof-gl).days+1;kk=ref[nm]["keys"];mg=ref[nm].get("mgmt","")
    rate=chg/bil if bil else 0;irf=bil/(kk*days) if kk and days else 0;profit=net-rc
    proj=int(bil/days*365) if days else 0;pb=(kk*350)/(proj*140) if proj else None
    g=grp.setdefault(mg,dict(bil=0,chg=0,net=0,rest=0,kd=0,keys=0,proj=0,n=0))
    g["bil"]+=bil;g["chg"]+=chg;g["net"]+=net;g["rest"]+=rc;g["kd"]+=kk*days;g["keys"]+=kk;g["proj"]+=proj;g["n"]+=1
    bf=GREY if idx%2==1 else None
    put(r,1,nm,al="left",fill=bf,italic=True,border=sides(1))
    put(r,2,gl,DATE,fill=bf,border=sides(2))
    put(r,3,f"=H{r}/({KL}{r}*{DL}{r})",PCT2,fill=bf,border=sides(3))
    put(r,4,f"=IF(H{r}=0,0,I{r}/H{r})",PCT0,fill=bf,border=sides(4))
    put(r,5,net,DOL0,al="right",fill=bf,border=sides(5))
    put(r,6,rc,DOL0,al="right",fill=bf,border=sides(6))
    put(r,7,f"=E{r}-F{r}",DOL0,al="right",fill=bf,border=sides(7))
    put(r,8,bil,INT0,fill=bf,border=sides(8))
    put(r,9,chg,INT0,fill=bf,border=sides(9))
    put(r,10,f"=E{r}/({KL}{r}*{DL}{r})",ACCD,fill=bf,border=sides(10))
    put(r,11,f"={KL}{r}*350",ACC0,fill=bf,border=sides(11))
    put(r,12,f"=INT(H{r}/{DL}{r}*365)",ACC0,fill=bf,border=sides(12))
    put(r,13,f'=IF(L{r}=0,"",K{r}/(L{r}*140))',ACC2,fill=bf,border=sides(13))
    put(r,14,mg,fill=bf,border=sides(14))
    for k,rcode in enumerate(rorder):put(r,RBASE+k,R.get(rcode,0),INT0,fill=bf,border=Border())
    put(r,KEYC,kk,INT0,fill=bf,border=Border());put(r,DAYC,f"=TODAY()-B{r}+1","0",fill=bf,border=Border())
    ws.row_dimensions[r].height=15.75

tr=r0+len(rows)
Tb=sum(g["bil"] for g in grp.values());Tc=sum(g["chg"] for g in grp.values());Tn=sum(g["net"] for g in grp.values());Trs=sum(g["rest"] for g in grp.values());Tkd=sum(g["kd"] for g in grp.values());Tk=sum(g["keys"] for g in grp.values());Tp=sum(g["proj"] for g in grp.values())
def totb(j):return sides(j,top=thin,bottom=med)
def tput(col,val,fmt=None,al="center",italic=False):
    c=ws.cell(tr,col,val);c.font=Font(name=FN,size=16,bold=True,italic=italic);c.alignment=Alignment(horizontal=al,vertical="center")
    if fmt:c.number_format=fmt
    c.fill=PatternFill("solid",fgColor=GREY);c.border=totb(col);return c
tput(1,"Total: "+str(len(rows)),al="left",italic=True);tput(2,"")
tput(3,f"=H{tr}/SUMPRODUCT({KL}{r0}:{KL}{tr-1},{DL}{r0}:{DL}{tr-1})",PCT2)
tput(4,f"=IF(H{tr}=0,0,I{tr}/H{tr})",PCTd)
tput(5,f"=SUM(E{r0}:E{tr-1})",DOL0,al="right");tput(6,f"=SUM(F{r0}:F{tr-1})",DOL0,al="right")
tput(7,f"=E{tr}-F{tr}",DOL0,al="right")
tput(8,f"=SUM(H{r0}:H{tr-1})",ACC0);tput(9,f"=SUM(I{r0}:I{tr-1})",ACC0)
tput(10,f"=E{tr}/SUMPRODUCT({KL}{r0}:{KL}{tr-1},{DL}{r0}:{DL}{tr-1})",ACCD)
tput(11,f"={KL}{tr}*350",ACC0);tput(12,f"=SUM(L{r0}:L{tr-1})",ACC0)
tput(13,f"=K{tr}/(L{tr}*140)",ACC2)
tput(14,"")
for k in range(len(rorder)):
    cc=ws.cell(tr,RBASE+k,f"=SUM({get_column_letter(RBASE+k)}{r0}:{get_column_letter(RBASE+k)}{tr-1})");cc.font=Font(name=FN,size=16,bold=True);cc.number_format=ACC0
cc=ws.cell(tr,KEYC,f"=SUM({KL}{r0}:{KL}{tr-1})");cc.font=Font(name=FN,size=16,bold=True)
ws.row_dimensions[tr].height=21.75

# ---- management-company rollup (own boxed table) ----
NR=f"$N${r0}:$N${tr-1}"; KR=f"${KL}${r0}:${KL}${tr-1}"; DR=f"${DL}${r0}:${DL}${tr-1}"
groups=sorted(grp); g0=tr+2; lastg=g0+len(groups)-1
for gi,gname in enumerate(groups):
    r=g0+gi; bf=GREY if gi%2==1 else None
    def gb(j):return sides(j, top=(med if gi==0 else None), bottom=(med if gi==len(groups)-1 else None))
    st=grp[gname]; kd=st["kd"]
    grate=st["chg"]/st["bil"] if st["bil"] else 0; girf=st["bil"]/kd if kd else 0
    gprof=st["net"]-st["rest"]; gpb=(st["keys"]*350)/(st["proj"]*140) if st["proj"] else None
    put(r,1,gname,al="left",fill=bf,italic=True,border=gb(1))
    put(r,2,f'=COUNTIF($N${r0}:$N${tr-1},A{r})&" Hotels"',fill=bf,border=gb(2))
    put(r,3,f'=H{r}/SUMPRODUCT(({NR}=A{r})*{KR}*{DR})',PCT2,fill=bf,border=gb(3))
    put(r,4,f"=IF(H{r}=0,0,I{r}/H{r})",PCTd,fill=bf,border=gb(4))
    put(r,5,f"=SUMIF({NR},A{r},E{r0}:E{tr-1})",DOL0,al="right",fill=bf,border=gb(5))
    put(r,6,f"=SUMIF({NR},A{r},F{r0}:F{tr-1})",DOL0,al="right",fill=bf,border=gb(6))
    put(r,7,f"=E{r}-F{r}",DOL0,al="right",fill=bf,border=gb(7))
    put(r,8,f"=SUMIF({NR},A{r},H{r0}:H{tr-1})",INT0,fill=bf,border=gb(8))
    put(r,9,f"=SUMIF({NR},A{r},I{r0}:I{tr-1})",INT0,fill=bf,border=gb(9))
    put(r,10,f'=E{r}/SUMPRODUCT(({NR}=A{r})*{KR}*{DR})',ACCD,fill=bf,border=gb(10))
    put(r,11,f"=SUMIF({NR},A{r},{KR})*350",ACC0,fill=bf,border=gb(11))
    put(r,12,f"=SUMIF({NR},A{r},L{r0}:L{tr-1})",ACC0,fill=bf,border=gb(12))
    put(r,13,f'=IF(L{r}=0,"",K{r}/(L{r}*140))',ACC2,fill=bf,border=gb(13))
    put(r,14,"",fill=bf,border=gb(14))
    ws.row_dimensions[r].height=15.75

# ---- widths (match master visible columns) ----
W={"A":67.86,"B":15.14,"C":12.14,"D":9.14,"E":13.14,"F":14.43,"G":13.14,"H":12.71,"I":11.14,"J":10.14,"K":16.0,"L":17.29,"M":22.29,"N":15.86}
for col,w in W.items():ws.column_dimensions[col].width=w
for k in range(len(rorder)):
    cd=ws.column_dimensions[get_column_letter(RBASE+k)];cd.width=13;cd.hidden=True;cd.outlineLevel=1
for L in (KL,DL):
    cd=ws.column_dimensions[L];cd.width=10;cd.hidden=True;cd.outlineLevel=1
ws.sheet_properties.outlinePr.summaryRight=True
ws.freeze_panes="C2"
ws.print_area="A1:N"+str(lastg)
ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.page_margins.left=ws.page_margins.right=ws.page_margins.top=ws.page_margins.bottom=0.2
# ---- conditional formatting: same rules on data rows, total, and rollup ----
def CF(colL,op,val,rgb):
    rule=CellIsRule(operator=op,formula=[val],font=Font(bold=True,color=rgb))
    for rng in (f"{colL}{r0}:{colL}{tr}", f"{colL}{g0}:{colL}{lastg}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator=op,formula=[val],font=Font(bold=True,color=rgb)))
CF("C","lessThan","0.006","C00000")   # Incident/Av Room < 0.60% -> red
CF("D","lessThan","0.75","C00000")    # Charge Rate < 75% -> red
CF("G","lessThan","0","C00000")       # Profit < 0 -> red
CF("M","lessThan","1.26","00B050")    # Payback < 1.26 -> green
wb.save(out);print("saved",out,"| groups",",".join(groups),"| total",tr,"| rollup",g0,"-",lastg)
