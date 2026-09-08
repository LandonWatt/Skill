#!/usr/bin/env python3
"""Render the email-format Excel to a cropped PNG for embedding in the Word doc.
Usage: python render_table.py <Email.xlsx> <asof YYYY-MM-DD> <out.png>
- Freezes Days-Since-Go-Live to the report date (picture is a point-in-time snapshot).
- Fills in static values for the hotel table AND the management-company rollup so no
  LibreOffice recalc is needed.
- Aliases "Aptos Narrow"/"Aptos" to installed narrow fonts so the size-16 Total row
  never overflows to ###.
Requires: openpyxl, LibreOffice (soffice), pdftoppm, PIL.
"""
import sys, os, subprocess, tempfile
from datetime import date
from openpyxl import load_workbook

src, asof_s, out = sys.argv[1], sys.argv[2], sys.argv[3]
asof = date.fromisoformat(asof_s)
wb = load_workbook(src); ws = wb.active

# locate total row and key columns by header text
tr = next(r for r in range(1, ws.max_row+1)
          if isinstance(ws.cell(r,1).value, str) and ws.cell(r,1).value.startswith("Total"))
KEYC = DAYC = MGMTC = None
for c in range(1, ws.max_column+1):
    h = ws.cell(1,c).value
    if h == "Keys": KEYC = c
    if h == "Days Since Go-Live": DAYC = c
    if h == "Management Company": MGMTC = c
r0 = 2
def num(r,c):
    v = ws.cell(r,c).value
    return v if isinstance(v,(int,float)) else 0

grp = {}            # mgmt -> aggregates
sE=sF=sH=sI=sL=sKD=sK=0
for r in range(r0, tr):
    b = ws.cell(r,2).value
    bd = b.date() if hasattr(b,"date") else b
    days = (asof - bd).days + 1 if hasattr(bd,"year") else num(r,DAYC)
    keys = num(r,KEYC)
    bil, chg, net, rest = num(r,8), num(r,9), num(r,5), num(r,6)
    mg = ws.cell(r,MGMTC).value if MGMTC else None
    kd = keys*days
    ws.cell(r,3).value  = bil/kd if kd else 0
    ws.cell(r,4).value  = chg/bil if bil else 0
    ws.cell(r,7).value  = net - rest
    ws.cell(r,10).value = net/kd if kd else 0
    ws.cell(r,11).value = keys*350
    proj = int(bil/days*365) if days else 0
    ws.cell(r,12).value = proj
    pb = (keys*350)/(proj*140) if proj else None
    ws.cell(r,13).value = pb if pb is not None else ""
    ws.cell(r,DAYC).value = days
    sE+=net; sF+=rest; sH+=bil; sI+=chg; sL+=proj; sKD+=kd; sK+=keys
    if mg:
        g = grp.setdefault(mg, dict(bil=0,chg=0,net=0,rest=0,kd=0,keys=0,proj=0,n=0))
        g["bil"]+=bil; g["chg"]+=chg; g["net"]+=net; g["rest"]+=rest; g["kd"]+=kd; g["keys"]+=keys; g["proj"]+=proj; g["n"]+=1
# grand total row
ws.cell(tr,3).value  = sH/sKD if sKD else 0
ws.cell(tr,4).value  = sI/sH if sH else 0
ws.cell(tr,5).value  = sE; ws.cell(tr,6).value = sF; ws.cell(tr,7).value = sE - sF
ws.cell(tr,8).value  = sH; ws.cell(tr,9).value = sI
ws.cell(tr,10).value = sE/sKD if sKD else 0
ws.cell(tr,11).value = sK*350; ws.cell(tr,12).value = sL
ws.cell(tr,13).value = (sK*350)/(sL*140) if sL else ""
# management-company rollup rows (col A == a group name)
lastrow = tr
for r in range(tr+1, ws.max_row+1):
    nm = ws.cell(r,1).value
    if not (isinstance(nm,str) and nm in grp): continue
    g = grp[nm]; kd = g["kd"]
    ws.cell(r,2).value = f'{g["n"]} Hotels'
    ws.cell(r,3).value  = g["bil"]/kd if kd else 0
    ws.cell(r,4).value  = g["chg"]/g["bil"] if g["bil"] else 0
    ws.cell(r,5).value  = g["net"]; ws.cell(r,6).value = g["rest"]; ws.cell(r,7).value = g["net"]-g["rest"]
    ws.cell(r,8).value  = g["bil"]; ws.cell(r,9).value = g["chg"]
    ws.cell(r,10).value = g["net"]/kd if kd else 0
    ws.cell(r,11).value = g["keys"]*350; ws.cell(r,12).value = g["proj"]
    ws.cell(r,13).value = (g["keys"]*350)/(g["proj"]*140) if g["proj"] else ""
    lastrow = r

# page setup: single landscape page, tight margins
ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
ws.page_setup.scale=None; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.paperSize=8
ws.print_area=f"A1:N{lastrow}"; ws.print_options.horizontalCentered=True
ws.page_margins.left=ws.page_margins.right=ws.page_margins.top=ws.page_margins.bottom=0.2

tmp = tempfile.mkdtemp(prefix="resttbl_")
xlsx = os.path.join(tmp,"render.xlsx"); wb.save(xlsx)

conf = os.path.join(tmp,"fonts.conf")
open(conf,"w").write("""<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "fonts.dtd">
<fontconfig>
  <dir>/usr/share/fonts</dir><cachedir>%s/fc</cachedir>
  <match target="pattern"><test name="family"><string>Aptos Narrow</string></test><edit name="family" mode="assign" binding="strong"><string>Liberation Sans Narrow</string></edit></match>
  <match target="pattern"><test name="family"><string>Aptos</string></test><edit name="family" mode="assign" binding="strong"><string>Liberation Sans</string></edit></match>
</fontconfig>""" % tmp)
env = dict(os.environ, FONTCONFIG_FILE=conf, HOME=tmp)
subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",tmp,xlsx],
               env=env, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
pdf = os.path.join(tmp,"render.pdf")
subprocess.run(["pdftoppm","-png","-r","200",pdf,os.path.join(tmp,"pg")], check=True)

from PIL import Image, ImageChops
im = Image.open(os.path.join(tmp,"pg-1.png")).convert("RGB")
bbox = ImageChops.difference(im, Image.new("RGB",im.size,(255,255,255))).getbbox()
l,t,r,b = bbox; p=12
im.crop((max(0,l-p),max(0,t-p),min(im.size[0],r+p),min(im.size[1],b+p))).save(out)
print("saved", out)
